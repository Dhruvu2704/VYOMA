from pathlib import Path
from typing import Any, Mapping

from docx import Document


def generate_word(
    final_verdict: Mapping[str, Any],
    output_dir: str | Path = "outputs",
) -> dict[str, str]:
    """
    Generate a safety review memo from the shared FinalVerdict structure.

    The function consumes the final verdict produced by the project
    pipeline and creates a Word document for human review.
    """

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    permit_id = final_verdict["permit_id"]
    file_path = output_path / f"{permit_id}_memo.docx"

    document = Document()

    # Title
    document.add_heading("KAVACH Safety Review Memo", level=0)

    # Permit information
    document.add_heading("Permit Information", level=1)
    document.add_paragraph(f"Permit ID: {permit_id}")

    # Decision information
    document.add_heading("Safety Decision", level=1)
    document.add_paragraph(
        f"Final Decision: {final_verdict['final_decision']}"
    )
    document.add_paragraph(
        f"Rule Result: {final_verdict['rule_result']}"
    )
    document.add_paragraph(
        f"LLM Result: {final_verdict['llm_result']}"
    )
    document.add_paragraph(
        f"Agreement: {final_verdict['agreement']}"
    )
    document.add_paragraph(
        f"Human Review Required: "
        f"{final_verdict['requires_human_review']}"
    )

    # Explanation
    document.add_heading("Explanation", level=1)
    document.add_paragraph(final_verdict["explanation"])

    # Audit information
    document.add_heading("Audit Information", level=1)
    document.add_paragraph(
        f"Generated At: {final_verdict['generated_at']}"
    )

    audit_ref = final_verdict.get("audit_ref")
    document.add_paragraph(
        f"Audit Reference: {audit_ref if audit_ref else 'N/A'}"
    )

    # Save document
    document.save(file_path)

    return {
        "type": "WORD_MEMO",
        "path": str(file_path),
    }
