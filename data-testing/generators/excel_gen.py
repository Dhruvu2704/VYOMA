from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def generate_excel(
    rule_verdict: Mapping[str, Any] | Sequence[Mapping[str, Any]],
    output_dir: str | Path = "outputs",
) -> dict[str, str]:
    """
    Generate a conflict matrix from RuleVerdict data.

    The function accepts either one RuleVerdict mapping or a sequence
    of RuleVerdict mappings and creates an Excel evidence file.
    """

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    if isinstance(rule_verdict, Mapping):
        verdicts = [rule_verdict]
    else:
        verdicts = list(rule_verdict)

    file_path = output_path / "permit_conflict_matrix.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Conflict Matrix"

    headers = [
        "Permit ID",
        "Conflicting Permit IDs",
        "Rules Triggered",
        "Rule Result",
        "Confidence",
        "Explanation",
    ]

    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for verdict in verdicts:
        conflicting_permits = verdict.get("conflicting_permit_ids", [])
        rules_triggered = verdict.get("rules_triggered", [])

        if not isinstance(conflicting_permits, list):
            conflicting_permits = [str(conflicting_permits)]

        if not isinstance(rules_triggered, list):
            rules_triggered = [str(rules_triggered)]

        worksheet.append(
            [
                verdict["permit_id"],
                ", ".join(conflicting_permits),
                ", ".join(rules_triggered),
                verdict["rule_result"],
                verdict["confidence"],
                verdict["explanation"],
            ]
        )

    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        worksheet.column_dimensions[column_letter].width = min(
            max_length + 2,
            60,
        )

    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    workbook.save(file_path)

    return {
        "type": "EXCEL_CONFLICT_MATRIX",
        "path": str(file_path),
    }