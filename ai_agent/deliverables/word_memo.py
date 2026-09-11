"""Word memo generator for KAVACH safety review deliverables.

Adapted from Role 6 data-testing/generators/word_gen.py.
Generates a .docx safety review memo from FinalVerdict data.
"""

from pathlib import Path
from typing import Any, Mapping


def generate_word(
    final_verdict: Mapping[str, Any],
    output_dir: str | Path = "outputs",
) -> dict[str, str]:
    """Generate a safety review memo from the shared FinalVerdict structure.

    Args:
        final_verdict: A dictionary containing permit_id, final_decision,
            rule_result, llm_result, agreement, requires_human_review,
            explanation, and optional generated_at / audit_ref.
        output_dir: Directory to write the generated .docx file into.

    Returns:
        A dict with keys "type" and "path".
    """

    from docx import Document

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    permit_id = final_verdict["permit_id"]
    file_path = output_path / f"{permit_id}_memo.docx"

    document = Document()

    document.add_heading("KAVACH Safety Review Memo", level=0)

    document.add_heading("Permit Information", level=1)
    document.add_paragraph(f"Permit ID: {permit_id}")

    document.add_heading("Safety Decision", level=1)
    document.add_paragraph(
        f"Final Decision: {final_verdict['final_decision']}"
    )
    document.add_paragraph(
        f"Rule Result: {final_verdict['rule_result']}"
    )
    document.add_paragraph(
        f"LLM Result: {final_verdict['llm_result']}"
    )
    document.add_paragraph(
        f"Agreement: {final_verdict['agreement']}"
    )
    document.add_paragraph(
        f"Human Review Required: "
        f"{final_verdict['requires_human_review']}"
    )

    document.add_heading("Explanation", level=1)
    document.add_paragraph(final_verdict["explanation"])

    document.add_heading("Audit Information", level=1)
    generated_at = final_verdict.get("generated_at", "")
    document.add_paragraph(f"Generated At: {generated_at}")

    audit_ref = final_verdict.get("audit_ref")
    document.add_paragraph(
        f"Audit Reference: {audit_ref if audit_ref else 'N/A'}"
    )

    document.save(file_path)

    return {
        "type": "WORD_MEMO",
        "path": str(file_path),
    }
